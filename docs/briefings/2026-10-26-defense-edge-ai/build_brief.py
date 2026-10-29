"""Build the visual Defense scenario brief and editable/text versions."""
from pathlib import Path
import json, html, io
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
HERE=Path(__file__).resolve().parent
D=json.loads((HERE/'content.json').read_text())
for name,file in [('Brief','LiberationSans-Regular.ttf'),('Brief-Bold','LiberationSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/liberation/'+file))
pdfmetrics.registerFontFamily('Brief',normal='Brief',bold='Brief-Bold')
W,H=612,792; M=42; CW=528
INK=HexColor('#163D33'); TEXT=HexColor('#30443A'); MUTED=HexColor('#64776B'); LINE=HexColor('#D1DDD4')
BODY=ParagraphStyle('body',fontName='Brief',fontSize=11,leading=15,textColor=TEXT)
SMALL=ParagraphStyle('small',fontName='Brief',fontSize=10,leading=14,textColor=MUTED)
CAP=ParagraphStyle('cap',fontName='Brief',fontSize=8,leading=10.5,textColor=MUTED)
HEAD=ParagraphStyle('head',fontName='Brief-Bold',fontSize=17,leading=21,textColor=INK)
BOUNDS=[]
def p(c,text,x,y,width,style=BODY):
    q=Paragraph(html.escape(text),style);_,h=q.wrap(width,900);q.drawOn(c,x,y-h);return y-h

def bullets(c,items,x,y,width,style=BODY):
    for text in items:
        c.setFillColor(INK);c.circle(x+2,y-6,1.5,fill=1,stroke=0)
        y=p(c,text,x+13,y,width-13,style)-8
    return y

def image(c,name,x,y,width):
    path=HERE/'assets'/name
    im=Image.open(path).convert('RGB')
    height=width*im.height/im.width
    b=io.BytesIO(); im.save(b,'JPEG',quality=95,optimize=True);b.seek(0)
    c.drawImage(ImageReader(b),x,y-height,width=width,height=height)
    return y-height

def base(c,page,total):
    c.setFillColor(INK);c.rect(0,H-7,W,7,fill=1,stroke=0)
    c.setFont('Brief-Bold',9);c.drawString(M,H-31,'CONNECT TECH')
    c.setFillColor(MUTED);c.setFont('Brief',8);c.drawRightString(W-M,H-31,'DEFENSE / SCENARIOS')
    c.setStrokeColor(LINE);c.line(M,38,W-M,38)
    c.setFont('Brief',7.5);c.drawString(M,25,'26 October 2026  |  Scenario concepts and illustrative imagery')
    c.drawRightString(W-M,25,f'{page} / {total}')

def title(c,text,sub):
    c.setFillColor(INK);c.setFont('Brief-Bold',27);c.drawString(M,725,text)
    c.setFillColor(MUTED);c.setFont('Brief',12);c.drawString(M,704,sub)
    return 680

def finish(c,page,y):
    if y<51:raise RuntimeError(f'Page {page} too low: {y}')
    BOUNDS.append({'page':page,'content_bottom':round(y,1),'footer_top':38});c.showPage()

def scenario_page(c,i):
    s=D['scenarios'][i];base(c,i+1,4)
    if i==0:
        y=title(c,D['title'],D['subtitle'])
        y=p(c,D['intro'],M,y,CW)-20
        y=p(c,'01  '+s['title'],M,y,CW,HEAD)-8
        y=p(c,s['subtitle'],M,y,CW,SMALL)-15
    else:
        y=title(c,f'0{i+1}  '+s['title'],s['subtitle'])
    y=image(c,s['image'],M,y,CW)
    y=p(c,s['caption'],M,y-7,CW,CAP)-18
    y=bullets(c,s['bullets'],M,y,CW)-3
    if i<2:
        c.setStrokeColor(LINE);c.line(M,y,W-M,y)
        y=p(c,s['example'],M,y-13,CW,SMALL)
    else:
        y=p(c,'How this fits into VSS',M,y-7,CW,HEAD)-8
        y=p(c,D['vss_intro'],M,y,CW)-9
        y=bullets(c,D['vss_bullets'],M,y,CW,SMALL)
    finish(c,i+1,y)

def inspection_page(c):
    base(c,4,4);y=title(c,'Inspection and border security','See the activity. Bring the relevant images forward.')
    for i,s in enumerate(D['scenarios'][3:],4):
        y=p(c,f'0{i}  '+s['title'],M,y,CW,HEAD)-12
        bottom=image(c,s['image'],M,y,250)
        caption_bottom=p(c,s['caption'],M,bottom-7,250,CAP)
        text_bottom=bullets(c,s['bullets'],M+269,y,CW-269)
        y=min(caption_bottom,text_bottom)-29
        if i==4:
            c.setStrokeColor(LINE);c.line(M,y+12,W-M,y+12)
    finish(c,4,y)

def text_pdf():
    c=canvas.Canvas(str(HERE/'Connect-Tech-Defense-Executive-Brief-Text-Only.pdf'),pagesize=(W,H))
    c.setTitle(D['title']);c.setAuthor('Connect Tech');base(c,1,1)
    y=title(c,D['title'],D['subtitle']);y=p(c,D['intro'],M,y,CW,SMALL)-13
    for i,s in enumerate(D['scenarios'],1):
        y=p(c,f'0{i}  '+s['title'],M,y,CW,HEAD)-6
        y=bullets(c,s['bullets'],M,y,CW,SMALL)-5
    y=p(c,'How this fits into VSS',M,y,CW,HEAD)-6
    y=p(c,D['vss_intro']+' Watch the views, inspect findings and ask questions with supporting images or clips.',M,y,CW,SMALL)
    if y<51:raise RuntimeError(f'Text-only overflow: {y}')
    c.save()

def word():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.58);sec.left_margin=sec.right_margin=Inches(.58)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Liberation Sans';st.font.color.rgb=RGBColor.from_string('30443A')
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.05
    for name,size in [('Title',27),('Heading 1',25),('Heading 2',17)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('163D33')
    cap=doc.styles['Caption'];cap.font.size=Pt(8);cap.font.italic=False;cap.font.bold=False;cap.font.color.rgb=RGBColor.from_string('64776B')
    sec.header.paragraphs[0].text='CONNECT TECH / DEFENSE SCENARIOS';sec.header.paragraphs[0].style='Caption'
    sec.footer.paragraphs[0].text='26 October 2026 | Scenario concepts and illustrative imagery';sec.footer.paragraphs[0].style='Caption'
    for i,s in enumerate(D['scenarios'][:3]):
        if i:doc.add_page_break()
        if i==0:
            doc.add_heading(D['title'],0);doc.add_paragraph(D['subtitle']);doc.add_paragraph(D['intro'])
        doc.add_heading(f'0{i+1}  '+s['title'],1 if i else 2)
        doc.add_paragraph(s['subtitle'])
        doc.add_picture(str(HERE/'assets'/s['image']),width=Inches(7.3))
        doc.add_paragraph(s['caption'],'Caption')
        for text in s['bullets']:doc.add_paragraph(text,'List Bullet')
        if i<2:doc.add_paragraph(s['example'])
        else:
            doc.add_heading('How this fits into VSS',2);doc.add_paragraph(D['vss_intro'])
            for text in D['vss_bullets']:doc.add_paragraph(text,'List Bullet')
    doc.add_page_break();doc.add_heading('Inspection and border security',1)
    doc.add_paragraph('See the activity. Bring the relevant images forward.')
    for i,s in enumerate(D['scenarios'][3:],4):
        doc.add_heading(f'0{i}  '+s['title'],2)
        t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(3.5);t.columns[1].width=Inches(3.8)
        a,b=t.rows[0].cells
        a.paragraphs[0].add_run().add_picture(str(HERE/'assets'/s['image']),width=Inches(3.4))
        a.add_paragraph(s['caption'],'Caption')
        for j,text in enumerate(s['bullets']):
            pp=b.paragraphs[0] if j==0 else b.add_paragraph();pp.style='List Bullet';pp.add_run(text)
    doc.core_properties.title=D['title'];doc.core_properties.author='Connect Tech'
    doc.save(str(HERE/'Connect-Tech-Defense-Executive-Brief.docx'))

if __name__=='__main__':
    c=canvas.Canvas(str(HERE/'Connect-Tech-Defense-Executive-Brief.pdf'),pagesize=(W,H))
    c.setTitle('Connect Tech | '+D['title']);c.setAuthor('Connect Tech')
    for i in range(3):scenario_page(c,i)
    inspection_page(c);c.save();text_pdf();word();print(json.dumps(BOUNDS))
